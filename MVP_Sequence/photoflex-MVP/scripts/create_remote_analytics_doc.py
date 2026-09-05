from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "docs/PhotoFlex_产品行为数据远程方案.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    if italic is not None:
        run.italic = italic


def style_paragraph(paragraph, size=10.5, color=(45, 55, 72), space_after=5, line=1.18):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2563A8")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    new_run.append(r_fonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color=(20, 51, 93))
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=10.5, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    style_paragraph(p)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    style_paragraph(p, space_after=3)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.05
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6FA")
    pPr.append(shd)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(35, 50, 70)
    return p


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
        set_cell_shading(hdr[i], "173F6B")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                set_run_font(r, size=9.5, bold=True, color=(255, 255, 255))
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], "F5F8FC")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    set_run_font(r, size=9, color=(45, 55, 72))
        
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top={"val": "single", "sz": 4, "color": "D8E1EC"}, bottom={"val": "single", "sz": 4, "color": "D8E1EC"}, left={"val": "single", "sz": 4, "color": "D8E1EC"}, right={"val": "single", "sz": 4, "color": "D8E1EC"})
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8, color=(120, 130, 145))
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8, color=(120, 130, 145))


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(45, 55, 72)

    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(5)
    r = title.add_run("PhotoFlex 产品行为数据远程方案")
    set_run_font(r, size=22, bold=True, color=(20, 51, 93))
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    r = sub.add_run("给 MVP 试用阶段的简单落地指南")
    set_run_font(r, size=11, color=(92, 106, 124))

    add_body(doc, "目标：让你在远程看到用户如何使用 PhotoFlex，从而判断产品哪里卡住、哪里有价值，同时不把用户的项目文件和原始照片上传到服务器。")

    add_heading(doc, "一、先说结论", 1)
    add_body(doc, "推荐采用“本地项目数据 + 远程行为数据”的组合：项目和照片继续保存在用户电脑的浏览器里；每当用户完成一个重要动作，前端只发送一条很小的匿名事件到 Supabase。你随后可以在 Supabase 中查询数据，或接一个简单的管理页面查看漏斗。")
    add_body(doc, "这样最适合当前 MVP：改动小、成本低、容易撤回，也不会因为加统计功能而重做整个产品。")

    add_heading(doc, "二、目前 MVP 是什么架构？", 1)
    add_body(doc, "当前项目是 React + TypeScript + Vite 的前端应用。简单说，Page.tsx、组件和业务模块都运行在用户浏览器中；项目数据主要通过 IndexedDB 保存在本机。它目前没有一个传统意义上长期运行的后端服务器。")
    make_table(doc, ["部分", "现在在哪里运行", "它负责什么"], [
        ("页面与交互", "用户浏览器", "显示页面、响应点击、切换页面"),
        ("项目与照片索引", "用户浏览器本地", "保存项目、照片引用、序列草稿"),
        ("产品行为统计（新增）", "浏览器 → 远程接口", "记录用户完成了哪些动作"),
    ], [1.4, 1.55, 3.9])

    add_heading(doc, "三、推荐的远程数据流", 1)
    add_code(doc, "用户操作页面\n    │\n    ├─ 项目数据：继续写入浏览器 IndexedDB\n    │\n    └─ 行为事件：HTTPS → Supabase Edge Function → Postgres → 你的查询/看板")
    add_body(doc, "这里的关键点是分开两种数据：项目数据是用户的工作内容，行为事件只是“用户做了什么”的记录。第一阶段不要上传原始照片，也不要上传本地绝对路径。")

    doc.add_page_break()
    add_heading(doc, "四、第一版应该记录哪些行为？", 1)
    make_table(doc, ["事件名", "什么时候记录", "你能回答的问题"], [
        ("app_opened", "打开应用", "有多少人真的开始使用？"),
        ("project_created", "新建项目成功", "用户能否顺利开始？"),
        ("folder_added", "加入照片目录成功", "用户是否走到了核心流程？"),
        ("scan_failed", "扫描失败", "哪一步最容易出错？"),
        ("photo_selected", "选择一张照片", "用户是否找到了想要的照片？"),
        ("photo_added_to_table", "照片加入工作表", "核心动作是否完成？"),
        ("sequence_created", "创建序列", "用户是否理解序列功能？"),
        ("sequence_saved", "保存序列成功", "用户是否完成了一次有价值的产出？"),
        ("save_failed", "保存失败", "是否有数据丢失风险？"),
    ], [1.65, 2.3, 2.9])
    add_body(doc, "建议每条事件都带上少量通用字段：event_name、anonymous_user_id、session_id、app_version、route、occurred_at，以及必要的数量字段（例如照片数量）。")
    add_code(doc, '{\n  "event_name": "sequence_saved",\n  "anonymous_user_id": "随机生成的 ID",\n  "session_id": "本次打开应用的 ID",\n  "app_version": "0.1.0",\n  "route": "/sequence",\n  "photo_count": 24,\n  "occurred_at": "2026-09-04T10:30:00Z"\n}')

    add_heading(doc, "五、明确不要收集什么", 1)
    add_bullet(doc, "不要上传原始照片、缩略图或照片内容。")
    add_bullet(doc, "不要上传本地绝对路径，例如 C:\\Users\\某人\\照片。")
    add_bullet(doc, "不要默认上传项目名称、客户姓名、备注等可能包含隐私的文字。")
    add_bullet(doc, "不要把事件统计当成用户数据备份；事件丢失时，产品也应该能正常使用。")

    doc.add_page_break()
    add_heading(doc, "六、前端需要怎么改？", 1)
    add_body(doc, "为了让代码容易理解，可以把统计功能做成一个很小的接口。页面不直接调用 Supabase，而是只调用 analytics.track。这样以后更换服务时，页面代码基本不用动。")
    add_code(doc, "// src/app/AppDependencies.ts\nexport type Analytics = {\n  track: (eventName: string, properties?: Record<string, unknown>) => void;\n};\n\n// 页面或业务模块中\nanalytics.track('sequence_saved', { photo_count: photos.length });")
    add_body(doc, "建议先做两个实现：LocalAnalytics（开发时只在控制台打印）和 RemoteAnalytics（正式试用时发送到远程接口）。远程发送失败时只记录错误，不阻塞用户保存序列。")

    add_heading(doc, "七、Supabase 端需要什么？", 1)
    add_body(doc, "Supabase 可以同时提供数据库和一个很小的服务器函数。服务器函数的作用是接收浏览器事件、做基本校验，再写入数据库。这样数据库密钥不会直接放在前端代码里。")
    add_code(doc, "-- product_events 表（示意）\ncreate table product_events (\n  id bigint generated always as identity primary key,\n  event_name text not null,\n  anonymous_user_id text not null,\n  session_id text,\n  app_version text,\n  route text,\n  properties jsonb,\n  occurred_at timestamptz not null default now()\n);")
    add_body(doc, "第一阶段只需要一个表和一个 track 接口。等你确认产品真的需要更多数据，再增加登录、权限、团队看板等功能。")

    add_heading(doc, "八、远程看到数据后，重点看什么？", 1)
    make_table(doc, ["看板/问题", "最简单的计算方式", "如果数字很低，先检查"], [
        ("启动率", "app_opened 人数", "用户是否能成功打开应用"),
        ("开始率", "project_created ÷ app_opened", "新建项目入口是否清楚"),
        ("核心流程完成率", "photo_added_to_table ÷ folder_added", "扫描、筛选、加入工作表是否卡住"),
        ("产出率", "sequence_saved ÷ project_created", "用户是否理解序列，以及保存是否稳定"),
        ("失败率", "失败事件 ÷ 对应成功事件", "错误提示、权限和文件兼容性"),
    ], [1.7, 2.1, 3.05])
    add_body(doc, "MVP 不需要一开始就做复杂图表。你甚至可以先在 Supabase 的表格查询里按日期和 event_name 分组，先回答三个问题：用户有没有走到核心流程？在哪一步大量退出？失败是否集中在某一种文件或操作？")

    add_heading(doc, "九、建议的实现顺序", 1)
    for item in [
        "先生成 anonymous_user_id 和 session_id，并确认它们不会包含姓名或路径。",
        "先接 LocalAnalytics，在本地操作时确认事件名称和字段正确。",
        "在 Supabase 建 product_events 表和 track 接口。",
        "接 RemoteAnalytics；网络失败时静默降级，不影响原有功能。",
        "用自己的电脑完成一遍流程，检查远程表里是否出现正确事件。",
        "邀请 3–5 位试用者，先观察一周，再决定是否需要更多事件。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "十、服务器方案怎么选？", 1)
    make_table(doc, ["方案", "适合谁", "优点", "代价"], [
        ("Supabase（推荐）", "想快速做 MVP 的个人开发者", "数据库、接口、查询页面集中；上手快", "需要学习一点 SQL 和函数"),
        ("Cloudflare Workers + D1", "愿意自己搭更多基础设施的人", "轻量、全球边缘运行", "概念更多，调试链路更长"),
        ("Vercel Functions + 数据库", "前端已经部署在 Vercel", "部署体验顺手", "数据库仍要另选，服务分散"),
        ("PostHog 等托管分析", "只想尽快看到产品漏斗", "事件、漏斗、看板现成", "数据结构和费用受平台约束"),
    ], [1.55, 1.8, 2.0, 1.5])
    add_body(doc, "当前建议：先用 Supabase 做一版最小闭环。如果你只想验证用户路径、不想写查询页面，也可以先选 PostHog 这类托管分析服务。")

    add_heading(doc, "十一、第一阶段的完成标准", 1)
    add_bullet(doc, "用户仍然可以离线打开项目、添加照片、创建和保存序列。")
    add_bullet(doc, "你能在远程表里看到每次试用的匿名事件。")
    add_bullet(doc, "你能按日期、事件名和版本筛选数据。")
    add_bullet(doc, "网络断开时，用户不会因为统计失败而无法工作。")
    add_bullet(doc, "事件里没有原始照片、绝对路径或明显的个人隐私。")

    add_heading(doc, "十二、参考资料", 1)
    p = doc.add_paragraph()
    r = p.add_run("Supabase 数据库：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, "官方文档", "https://supabase.com/docs/guides/database/overview")
    style_paragraph(p, size=9.5, space_after=2)
    p = doc.add_paragraph()
    r = p.add_run("Supabase Edge Functions：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, "官方文档", "https://supabase.com/docs/guides/functions")
    style_paragraph(p, size=9.5, space_after=2)
    p = doc.add_paragraph()
    r = p.add_run("Cloudflare D1：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, "官方文档", "https://developers.cloudflare.com/d1/get-started/")
    style_paragraph(p, size=9.5, space_after=2)
    p = doc.add_paragraph()
    r = p.add_run("Vercel Functions：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, "官方文档", "https://vercel.com/docs/functions")
    style_paragraph(p, size=9.5, space_after=2)
    p = doc.add_paragraph()
    r = p.add_run("PostHog 事件：")
    set_run_font(r, size=9.5)
    add_hyperlink(p, "事件文档", "https://github.com/PostHog/posthog.com/blob/master/contents/docs/data/events.mdx")
    style_paragraph(p, size=9.5, space_after=2)

    doc.core_properties.title = "PhotoFlex 产品行为数据远程方案"
    doc.core_properties.subject = "MVP 阶段的匿名产品行为数据采集与远程查看"
    doc.core_properties.author = "OpenAI"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
